import random
from urllib import request
from django.utils import timezone
from django.shortcuts import render,redirect
from django.contrib.auth import authenticate, login,logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Sum
from django.http import HttpResponse
from django.contrib import messages
from .models import Transactions,Budget
from django.db.models.functions import ExtractMonth
from django.utils import timezone
from datetime import timedelta
from .emails import send_welcome_email, send_otp_email
from financial_ai.services import calculate_financial_health
@login_required(login_url="/login/")
@login_required(login_url="/login/")
def mainpage(request):
    today = timezone.now()

    title = request.GET.get("search_text")
    date = request.GET.get("search_date")
    types = request.GET.get("search_types")
    category = request.GET.get("search_category")
    amount = request.GET.get("search_amount")
    see_all = request.GET.get("see_all")
    month = request.GET.get("search_month")

    # Selected month (current month by default)
    selected_month = int(month) if month else today.month

    # Dashboard Cards
    expense = Transactions.objects.filter(
        user=request.user,
        types="Expense",
        date__year=today.year,
        date__month=selected_month
    ).aggregate(total_expense=Sum("amount"))

    income = Transactions.objects.filter(
        user=request.user,
        types="Income",
        date__year=today.year,
        date__month=selected_month
    ).aggregate(total_income=Sum("amount"))

    number_of_transactions = Transactions.objects.filter(
        user=request.user,
        date__year=today.year,
        date__month=selected_month
    ).count()

    Balance = (
        (income["total_income"] or 0)
        - (expense["total_expense"] or 0)
    )

    if Balance < 0:
        Balance = 0

    # Pie Chart
    expenseData = list(
        Transactions.objects.filter(
            user=request.user,
            types="Expense",
            date__year=today.year,
            date__month=selected_month
        )
        .values("category")
        .annotate(total=Sum("amount"))
    )

    # Transaction List
    queryset = Transactions.objects.filter(user=request.user).order_by("-created_at")

    if title:
        queryset = queryset.filter(title__icontains=title)

    if date:
        queryset = queryset.filter(date=date)

    if types:
        queryset = queryset.filter(types=types)

    if category:
        queryset = queryset.filter(category=category)

    if amount:
        queryset = queryset.filter(amount=amount)

    if not see_all:
        queryset = queryset.filter(
            date__year=today.year,
            date__month=selected_month
        )

    queryset = queryset[:10] if see_all else queryset[:5]

    # Budget Tracker
    budgets = Budget.objects.filter(user=request.user)

    budget_data = []

    for budget in budgets:

        spent = (
            Transactions.objects.filter(
                user=request.user,
                category=budget.category,
                types="Expense",
                date__year=today.year,
                date__month=selected_month
            ).aggregate(total=Sum("amount"))["total"]
            or 0
        )

        percentage = (spent / budget.limit) * 100 if budget.limit else 0

        budget_data.append({
            "id": budget.id,
            "category": budget.category,
            "limit": budget.limit,
            "spent": spent,
            "percentage": min(percentage, 100),
            "over_budget": spent > budget.limit,
            "remaining": max(budget.limit - spent, 0),
            "exceeded": max(spent - budget.limit, 0),
        })

    # ==========================
    # Income vs Expense Bar Chart
    # ==========================

    income_queryset = (
        Transactions.objects.filter(
            user=request.user,
            types="Income",
            date__year=today.year
        )
        .annotate(month=ExtractMonth("date"))
        .values("month")
        .annotate(total=Sum("amount"))
        .order_by("month")
    )

    expense_queryset = (
        Transactions.objects.filter(
            user=request.user,
            types="Expense",
            date__year=today.year
        )
        .annotate(month=ExtractMonth("date"))
        .values("month")
        .annotate(total=Sum("amount"))
        .order_by("month")
    )

    months = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"
    ]

    income_data = [0] * 12
    expense_data = [0] * 12

    for item in income_queryset:
        income_data[item["month"] - 1] = float(item["total"])

    for item in expense_queryset:
        expense_data[item["month"] - 1] = float(item["total"])

    health_preview = calculate_financial_health(request.user)

    return render(
        request,
        "expenses/mainpage.html",
        {
            "expense": expense,
            "income": income,
            "Balance": Balance,
            "number_of_transactions": number_of_transactions,
            "queryset": queryset,
            "expenseData": expenseData,
            
            "budget_data": budget_data,

            # Bar Chart
            "months": months,
            "income_data": income_data,
            "expense_data": expense_data,
            "health_preview": health_preview,
        },
    )
def budget_tracker(request):
    if request.method=="POST":
        category=request.POST.get("category")
        limit=request.POST.get("limit")
        if category and limit:
            Budget.objects.update_or_create(
                user=request.user,
                category=category,
                defaults={"limit":limit}
            )
        else:
            messages.info("Please select a category and set a budget limit for it ")
        budget = Budget.objects.filter(user=request.user)
        return redirect("/")
    return render(request,"expenses/budget_tracker.html")

def update_budget(request,id):
    budget = Budget.objects.get(user=request.user,id=id)
    if request.method=="POST":
        
        limit=request.POST.get("limit")
        budget.limit = limit
        budget.save()
        messages.info(request,"Budget updated successfully")
        return redirect("/")
    return render(request,"expenses/update_budget.html",{"budget": budget})

@login_required(login_url="/login/")
def add_transaction(request):
    if request.method=="POST":
        data=request.POST
        title=data.get("title")
        amount=data.get("amount")
        types=data.get("types")
        category=data.get("category")
        date=data.get("date")
        if title and amount and types and category:

            transaction = {
                "user": request.user,
                "title": title,
                "amount": amount,
                "types": types,
                "category": category,
            }

            if date:
                transaction["date"] = date

            Transactions.objects.create(**transaction)
            messages.info(request,"Transaction added successfully")
            return redirect("/add_transaction/")

    return render(request,"expenses/add_transaction.html")

# Create your views here.


@login_required(login_url="/login/")
def transaction(request):
    transactions = Transactions.objects.filter(
        user=request.user
    )

    return render(
        request,
        'expenses/transactions.html',
        {'tasks': transactions}
    )

def delete_transaction(request,id):
    queryset = Transactions.objects.get(
        user=request.user,id=id
    )
    queryset.delete()
    return redirect("/transactions/")

def update_transaction(request,id):
    queryset=Transactions.objects.get(user=request.user,id=id)
    if request.method == "POST":
        data=request.POST
        title=data.get("title")
        amount=data.get("amount")
        types=data.get("types")
        category=data.get("category")
        date=data.get("date")
        queryset.title = title
        queryset.amount = amount
        queryset.types = types
        queryset.category = category
        if date:
            queryset.date = date
        queryset.save()
        messages.info(request,"Transaction updated successfully")
        return redirect("/transactions/")
    return render(request,"expenses/update_transaction.html",{"task": queryset})




def register_user(request):
    if request.method == "POST":
        data = request.POST
        first_name = data.get("first_name")
        last_name = data.get("last_name")
        username = data.get("username")
        email = data.get("email")
        password = data.get("password")

        if User.objects.filter(username=username).first():
            messages.info(request, "Username already exists")
            return redirect('/register/')

        # Create user but inactive until OTP verified
        user = User.objects.create_user(
            username=username,
            email=email,
            first_name=first_name,
            last_name=last_name,
            password=password
        )
        user.is_active = False
        user.save()
        # Generate 6-digit OTP
        otp_code = str(random.randint(100000, 999999))
        # Store in session
        request.session['pending_user_id'] = user.id
        request.session['otp_code'] = otp_code
        request.session['otp_created_at'] = timezone.now().isoformat()

        send_otp_email(email, first_name, otp_code)
        return redirect('/verify-otp/')
    return render(request, "expenses/register.html")


def verify_otp(request):
    pending_user_id = request.session.get('pending_user_id')

    if not pending_user_id:
        messages.error(request, "No pending verification. Please register again.")
        return redirect('/register/')

    if request.method == "POST":
        entered_otp = request.POST.get("otp")
        stored_otp = request.session.get('otp_code')
        created_at_str = request.session.get('otp_created_at')

        created_at = timezone.datetime.fromisoformat(created_at_str)
        if timezone.now() > created_at + timedelta(minutes=10):
            messages.error(request, "OTP expired. Please register again.")
            User.objects.filter(id=pending_user_id).delete()
            request.session.flush()
            return redirect('/register/')

        if entered_otp == stored_otp:
            user = User.objects.get(id=pending_user_id)
            user.is_active = True
            user.save()

            send_welcome_email(user.email, user.first_name)

            login(request, user, backend='django.contrib.auth.backends.ModelBackend') #Since there are two backends i.e 
            #MOdelBAckend and allauth backend ,#login() function gets confuse ,therefore we mentioned it explicilty

            del request.session['pending_user_id']
            del request.session['otp_code']
            del request.session['otp_created_at']

            messages.success(request, "Account verified successfully!")
            return redirect('/')
        else:
            messages.error(request, "Invalid OTP. Try again.")
            return redirect('/verify-otp/')

    return render(request, "expenses/verify_otp.html")


def resend_otp(request):
    pending_user_id = request.session.get('pending_user_id')
    if not pending_user_id:
        return redirect('/register/')

    user = User.objects.get(id=pending_user_id)
    otp_code = str(random.randint(100000, 999999))

    request.session['otp_code'] = otp_code
    request.session['otp_created_at'] = timezone.now().isoformat()

    send_otp_email(user.email, user.first_name, otp_code)
    messages.info(request, "A new OTP has been sent to your email.")
    return redirect('/verify-otp/')

def login_user(request):
    if request.method=="POST":
        data=request.POST
        username=data.get("username")
        password=data.get("password")

        if not User.objects.filter(username=username).exists():
            messages.error(request,"Invalid Username")
            return redirect('/login/')
        
        user=authenticate(username=username,password=password)

        if user is None:
            messages.error(request,"Invalid Password")
            return redirect('/login/')
        elif not user.is_active:
            messages.error(request, "Please verify your email first.")
            return redirect('/login/')
        else:
            login(request,user)
            return redirect("/")

    return render(request,"expenses/login.html")

def logout_user(request):
    logout(request)
    return redirect('/login/')

login_required(login_url="/login/")
def profile_view(request):
    if request.method == "POST":
        user = request.user
        first_name = request.POST.get("first_name", "").strip()
        last_name = request.POST.get("last_name", "").strip()
        email = request.POST.get("email", "").strip()
 
        if email and User.objects.filter(email=email).exclude(id=user.id).exists():
            messages.error(request, "That email is already in use by another account")
            return redirect("/profile/")
 
        user.first_name = first_name
        user.last_name = last_name
        user.email = email
        user.save()
        messages.info(request, "Profile updated successfully")
        return redirect("/profile/")
 
    return render(request, "expenses/profile.html")
 
 
@login_required(login_url="/login/")
def change_password_view(request):
    if request.method == "POST":
        user = request.user
        current_password = request.POST.get("current_password")
        new_password = request.POST.get("new_password")
        confirm_password = request.POST.get("confirm_password")
 
        if not user.check_password(current_password):
            messages.error(request, "Current password is incorrect")
            return redirect("/profile/")
 
        if not new_password or new_password != confirm_password:
            messages.error(request, "New passwords do not match")
            return redirect("/profile/")
 
        user.set_password(new_password)
        user.save()
        # Re-authenticate so the user isn't logged out after changing their password
        login(request, user,backend='django.contrib.auth.backends.ModelBackend')
        messages.info(request, "Password updated successfully")
        return redirect("/profile/")
 
    return redirect("/profile/")

from io import BytesIO
from datetime import timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

@login_required(login_url="/login/")
def export_transactions_csv(request):
    transactions = Transactions.objects.filter(user=request.user)

    preset = request.GET.get('preset', 'current_month')
    today = timezone.now().date()
    label = "export"

    if preset == 'today':
        transactions = transactions.filter(date=today)
        label = str(today)

    elif preset == 'yesterday':
        yesterday = today - timedelta(days=1)
        transactions = transactions.filter(date=yesterday)
        label = str(yesterday)

    elif preset == 'current_month':
        transactions = transactions.filter(date__year=today.year, date__month=today.month)
        label = today.strftime('%Y-%m')

    elif preset == 'last_month':
        first_of_this_month = today.replace(day=1)
        last_month_end = first_of_this_month - timedelta(days=1)
        transactions = transactions.filter(date__year=last_month_end.year, date__month=last_month_end.month)
        label = last_month_end.strftime('%Y-%m')

    elif preset == 'custom':
        custom_type = request.GET.get('custom_type')

        if custom_type == 'month':
            month_value = request.GET.get('month')  # "YYYY-MM" from <input type="month">
            if month_value:
                year, month = map(int, month_value.split('-'))
                transactions = transactions.filter(date__year=year, date__month=month)
                label = month_value

        elif custom_type == 'date':
            date_value = request.GET.get('date')
            if date_value:
                transactions = transactions.filter(date=date_value)
                label = date_value

        elif custom_type == 'range':
            start_date = request.GET.get('start')
            end_date = request.GET.get('end')
            if start_date:
                transactions = transactions.filter(date__gte=start_date)
            if end_date:
                transactions = transactions.filter(date__lte=end_date)
            label = f"{start_date or 'start'}_to_{end_date or 'end'}"

    transactions = transactions.order_by('date')

    filename = f"spendwise_{request.user.username}_{label}.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transactions"
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    alternate_fill = PatternFill("solid", fgColor="F8FAFC")
    border = Border(bottom=Side(style="thin", color="E2E8F0"))

    worksheet.append(['Date', 'Title', 'Type', 'Category', 'Amount', 'Created At'])
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number, txn in enumerate(transactions, start=2):
        worksheet.append([
            txn.date,
            txn.title,
            txn.types,
            txn.category,
            txn.amount,
            timezone.localtime(txn.created_at).replace(tzinfo=None),
        ])
        for cell in worksheet[row_number]:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if row_number % 2 == 0:
                cell.fill = alternate_fill
        worksheet.cell(row=row_number, column=1).number_format = "yyyy-mm-dd"
        worksheet.cell(row=row_number, column=5).number_format = '₹#,##0.00'
        worksheet.cell(row=row_number, column=6).number_format = "yyyy-mm-dd hh:mm"

    for column, width in {"A": 14, "B": 28, "C": 14, "D": 18, "E": 16, "F": 22}.items():
        worksheet.column_dimensions[column].width = width

    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 23

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
